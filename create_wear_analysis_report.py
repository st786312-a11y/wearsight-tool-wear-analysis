from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path("reports/tool_wear_analysis_progress_report.docx")


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft JhengHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.65)
section.bottom_margin = Cm(1.65)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

styles = doc.styles
styles["Normal"].font.name = "Microsoft JhengHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(5)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(4)
run = title.add_run("車刀磨耗分析系統進度報告")
run.font.name = "Microsoft JhengHei"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
run.font.size = Pt(21)
run.font.color.rgb = RGBColor(0, 0, 0)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(14)
run = subtitle.add_run("YOLO 偵測、Dify 工作流與 Ollama 專家回答整合")
run.font.name = "Microsoft JhengHei"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
run.font.size = Pt(10.5)
run.font.color.rgb = RGBColor(85, 85, 85)

def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Microsoft JhengHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    r.font.size = Pt(12.5)
    r.font.color.rgb = RGBColor(20, 57, 104)


heading("一、目的")
doc.add_paragraph(
    "建立可操作的車刀分析原型：使用者提供圖片後，系統完成刀具偵測、產生框選結果，並可依使用者問題提供文字化說明。"
)

heading("二、目前系統流程")
flow = doc.add_paragraph()
flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
flow.paragraph_format.space_after = Pt(7)
run = flow.add_run("圖片網址或上傳圖片  →  YOLO 分析 API  →  問題分流  →  Ollama Expert 或直接輸出")
run.bold = True
run.font.name = "Microsoft JhengHei"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
run.font.size = Pt(10.5)
run.font.color.rgb = RGBColor(20, 57, 104)

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
table.autofit = False
widths = [Cm(3.1), Cm(6.7), Cm(5.1)]
for i, text in enumerate(["模組", "目前功能", "狀態"]):
    cell = table.rows[0].cells[i]
    cell.width = widths[i]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cell, "173968")
    set_cell_text(cell, text, bold=True, color=(255, 255, 255), size=9.5)

rows = [
    ("YOLO 偵測", "呼叫 best.pt，輸出刀具類別、偵測信心度與框選圖片。", "已串接"),
    ("Dify 工作流", "接收 image_url 與選填 question，呼叫 /api/analyze，依問題是否存在分流。", "已發布"),
    ("Ollama Expert", "使用本機 qwen2.5:1.5b，接收使用者問題與 API 分析結果，輸出繁中說明。", "已測試"),
    ("磨耗率", "目前為示意資料；尚非由 YOLO 直接預測的正式磨耗率。", "待驗證"),
]
for row_i, values in enumerate(rows):
    row = table.add_row()
    for col_i, value in enumerate(values):
        cell = row.cells[col_i]
        cell.width = widths[col_i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if row_i % 2 == 1:
            shade(cell, "EEF4FA")
        set_cell_text(cell, value, bold=(col_i == 0), size=9.3)

heading("三、驗證結果")
for item in [
    "Dify 已能成功呼叫本機 FastAPI，原先圖片輸入格式造成的 400 錯誤已排除。",
    "有輸入問題時，流程會進入 Ollama Expert，並已取得繁體中文回覆。",
    "未輸入問題時，流程直接回傳 YOLO 分析 JSON，保留框選圖網址與偵測資訊。",
]:
    p = doc.add_paragraph(style="Normal")
    p.style = doc.styles["Normal"]
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.add_run("• ").bold = True
    p.add_run(item)

heading("四、重要說明")
p = doc.add_paragraph()
p.add_run("YOLO 偵測信心度不等於磨耗率預測信心度。 ").bold = True
p.add_run("YOLO 目前負責找出刀具位置與類別；正式磨耗率應由後續的相似案例檢索、Top-K 與加權 kNN 驗證後產生。")

heading("五、下一步")
for item in [
    "將歷史案例圖片與已知磨耗標籤建立成資料庫。",
    "導入 CLIP 或 DINOv2 的 Top-5 相似案例檢索，並排除查詢圖片自我比對。",
    "以 similarity-weighted kNN 計算預估磨耗率，使用 MAE 與 RMSE 驗證可行性。",
]:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.add_run("• ").bold = True
    p.add_run(item)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT.resolve())
